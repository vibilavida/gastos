#!/usr/bin/env python3
# noqa: S603
"""Lee el Excel de revisión y envía las filas directamente al Google Sheet via Apps Script.

Usage:
    python3 send_to_sheets.py <xlsx_path> <script_url>

El script lee las filas del Excel (ya corregidas por la usuaria), construye
el payload JSON y hace un POST al Apps Script con action=write_batch.
"""

import json
import sys
from pathlib import Path

import requests
from openpyxl import load_workbook


COLUMNS_TO_SEND = ["Date", "Paid by", "Income/expense", "Amount", "Category",
                   "Subcategory", "Common", "Quien se beneficia"]


def read_xlsx(xlsx_path: Path) -> list[dict]:
    wb = load_workbook(xlsx_path, data_only=True)
    ws = wb.active

    header = [cell.value for cell in ws[1]]
    rows = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not any(row):
            continue
        row_dict = dict(zip(header, row))
        # Solo enviamos las columnas necesarias
        filtered = {k: (str(v) if v is not None else "") for k, v in row_dict.items()
                    if k in COLUMNS_TO_SEND}
        if filtered.get("Date"):
            rows.append(filtered)
    return rows


def send_batch(rows: list[dict], script_url: str) -> dict:
    try:
        resp = requests.post(
            script_url,
            json={"action": "write_batch", "rows": rows},
            timeout=90,
            allow_redirects=True
        )
        resp.raise_for_status()
        return resp.json()
    except Exception as e:
        return {"status": "error", "message": str(e)}


def main() -> int:
    if len(sys.argv) != 3:
        print("Uso: python3 send_to_sheets.py <xlsx_path> <script_url>")
        return 1

    xlsx_path = Path(sys.argv[1])
    script_url = sys.argv[2]

    if not xlsx_path.exists():
        print(f"ERROR: no encuentro el archivo {xlsx_path}")
        return 1

    rows = read_xlsx(xlsx_path)
    if not rows:
        print("No hay filas para enviar.")
        return 0

    print(f"Enviando {len(rows)} filas a Google Sheets...")
    result = send_batch(rows, script_url)

    if result.get("status") == "ok":
        print(f"OK: {result.get('written', len(rows))} filas escritas en el Sheet.")
        return 0
    else:
        print(f"ERROR: {result.get('message', result)}")
        return 1


if __name__ == "__main__":
    sys.exit(main())
